"""Form access and participant assignments."""

import json
import csv
from datetime import datetime
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from sqlalchemy import func
from openpyxl import load_workbook

from app.api.deps import get_current_user
from app.api.permissions import allowed_form_participant_ids, participant_visible, require_any_project_permission
from app.core.permissions import IDENTITY_USERS_MANAGE, PARTICIPANTS_CREATE, RECORDS_WRITE
from app.core.time import utc_now
from app.db.session import get_db
from app.models.audit import AuditLog
from app.models.builder import BuilderTemplate
from app.models.form_assignment import ParticipantFormAssignment
from app.models.identity import User
from app.models.participants import Participant
from app.models.runtime_record import RuntimeRecord
from app.schemas.participants import ParticipantRead
from app.services.assignment_service import assignment_service
from app.services.form_assignment_service import upsert_assignment, visible_assignment
from app.services.participant_service import _to_read as participant_to_read
from app.services.participant_source_service import configured_source, eligible_participants

router = APIRouter()


class AssignmentWrite(BaseModel):
    participant_id: str
    responsible_user_id: str


class AssignmentRead(BaseModel):
    id: str
    template_id: str
    participant_id: str
    participant_name: str
    document_id: str | None = None
    external_code: str | None = None
    record_id: str | None = None
    responsible_user_id: str
    responsible_name: str
    status: str
    created_at: datetime


class MyFormAssignmentRead(BaseModel):
    template_id: str
    template_name: str
    template_status: str
    participant_id: str
    participant_name: str
    document_id: str | None = None
    assignment_status: str
    record_id: str | None = None


class BulkAssignmentPreview(BaseModel):
    selected: int
    to_assign: int
    already_assigned: int
    to_reassign: int
    protected: int
    applied: int = 0
    issues: list[str] = Field(default_factory=list)


class OpenParticipantCreate(BaseModel):
    full_name: str = Field(min_length=2, max_length=220)
    document_id: str | None = Field(default=None, max_length=80)
    external_code: str | None = Field(default=None, max_length=120)
    department: str | None = Field(default=None, max_length=120)
    municipality: str | None = Field(default=None, max_length=120)


class ParticipantPhoneUpdate(BaseModel):
    phone: str = Field(min_length=7, max_length=25, pattern=r"^\+?[0-9 ()-]+$")


def _template(db: Session, user: User, template_id: str) -> BuilderTemplate:
    template = db.get(BuilderTemplate, template_id)
    if not template or not assignment_service.user_has_project_access(db, user.id, template.project_id):
        raise HTTPException(status_code=404, detail="Formulario no encontrado")
    return template


def _read(db: Session, row: ParticipantFormAssignment) -> AssignmentRead:
    participant = db.get(Participant, row.participant_id)
    responsible = db.get(User, row.responsible_user_id)
    record = db.query(RuntimeRecord.id).filter(RuntimeRecord.template_id == row.template_id,
        RuntimeRecord.participant_id == row.participant_id,
        RuntimeRecord.status.notin_(["voided", "cancelled"])).order_by(RuntimeRecord.created_at.desc()).first()
    return AssignmentRead(id=row.id, template_id=row.template_id, participant_id=row.participant_id,
                          participant_name=participant.full_name if participant else "Participante eliminado",
                          document_id=participant.document_id if participant else None,
                          external_code=participant.external_code if participant else None,
                          record_id=record[0] if record else None,
                          responsible_user_id=row.responsible_user_id,
                          responsible_name=responsible.full_name if responsible else "Usuario eliminado",
                          status=row.status, created_at=row.created_at)


@router.get("/mine/{project_id}", response_model=list[MyFormAssignmentRead])
def my_form_assignments(project_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> list[MyFormAssignmentRead]:
    if not assignment_service.user_has_project_access(db, user.id, project_id):
        raise HTTPException(status_code=403, detail="Sin acceso al proyecto")
    require_any_project_permission(db, user.id, project_id, {RECORDS_WRITE})
    rows = db.query(ParticipantFormAssignment, BuilderTemplate, Participant).join(
        BuilderTemplate, BuilderTemplate.id == ParticipantFormAssignment.template_id
    ).join(Participant, Participant.id == ParticipantFormAssignment.participant_id).filter(
        ParticipantFormAssignment.project_id == project_id,
        ParticipantFormAssignment.responsible_user_id == user.id,
        Participant.status == "active",
        BuilderTemplate.status == "published",
    ).order_by(BuilderTemplate.name, Participant.full_name).all()
    result: list[MyFormAssignmentRead] = []
    for assignment, template, participant in rows:
        if not participant_visible(db, user.id, participant):
            continue
        result.append(MyFormAssignmentRead(
            template_id=template.id, template_name=template.name, template_status=template.status,
            participant_id=participant.id, participant_name=participant.full_name,
            document_id=participant.document_id, assignment_status=assignment.status,
            record_id=_read(db, assignment).record_id,
        ))
    return result


@router.get("/templates/{template_id}", response_model=list[AssignmentRead])
def list_assignments(template_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> list[AssignmentRead]:
    template = _template(db, user, template_id)
    allowed = allowed_form_participant_ids(db, user.id, template)
    query = db.query(ParticipantFormAssignment).filter_by(template_id=template_id)
    if allowed is not None:
        query = query.filter(ParticipantFormAssignment.participant_id.in_(allowed))
    return [_read(db, row) for row in query.order_by(ParticipantFormAssignment.created_at).all()]


@router.get("/templates/{template_id}/candidates")
def assignment_candidates(template_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> list[dict[str, str | None]]:
    template = _template(db, user, template_id)
    require_any_project_permission(db, user.id, template.project_id, {IDENTITY_USERS_MANAGE})
    return [{"id": row.id, "full_name": row.full_name, "document_id": row.document_id,
             "external_code": row.external_code} for row in eligible_participants(db, template, enforce_assignments=False)]


def _excel_assignment_keys(filename: str, content: bytes) -> tuple[str, list[str]]:
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="El archivo supera 5 MB")
    try:
        if filename.lower().endswith(".xlsx"):
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
        elif filename.lower().endswith(".csv"):
            rows = csv.reader(StringIO(content.decode("utf-8-sig")))
        else:
            raise HTTPException(status_code=422, detail="Sube un archivo .xlsx o .csv")
        headers = [str(value or "").strip().casefold() for value in next(rows, [])]
        document_headers = {"documento", "document_id", "cedula", "cédula", "numero_documento"}
        code_headers = {"codigo", "código", "external_code", "codigo_participante"}
        document_columns = [index for index, header in enumerate(headers) if header in document_headers]
        code_columns = [index for index, header in enumerate(headers) if header in code_headers]
        if len(document_columns) + len(code_columns) != 1:
            raise HTTPException(status_code=422, detail="Usa una sola columna de identificación: documento o codigo")
        index = (document_columns or code_columns)[0]
        field = "document_id" if document_columns else "external_code"
        keys = []
        for row_number, row in enumerate(rows, 2):
            if row_number > 5001:
                raise HTTPException(status_code=422, detail="El archivo supera 5000 participantes")
            value = str(row[index]).strip() if index < len(row) and row[index] is not None else ""
            if value:
                keys.append(value)
        if not keys:
            raise HTTPException(status_code=422, detail="El archivo no contiene documentos o códigos")
        return field, keys
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"No fue posible leer el archivo: {exc}") from exc


@router.post("/templates/{template_id}/bulk-assign", response_model=BulkAssignmentPreview)
async def bulk_assign_participants(
    template_id: str,
    responsible_user_id: str = Form(...),
    mode: str = Form(...),
    group_name: str = Form(default=""),
    preview_only: bool = Form(default=True),
    upload: UploadFile | None = File(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> BulkAssignmentPreview:
    template = _template(db, user, template_id)
    require_any_project_permission(db, user.id, template.project_id, {IDENTITY_USERS_MANAGE})
    if configured_source(template).access_mode == "legacy":
        raise HTTPException(status_code=422, detail="Selecciona primero el modo abierto o cerrado del formulario")
    if not assignment_service.user_has_project_access(db, responsible_user_id, template.project_id):
        raise HTTPException(status_code=422, detail="El responsable no tiene acceso al proyecto")
    candidates = eligible_participants(db, template, enforce_assignments=False)
    issues: list[str] = []
    if mode == "group":
        wanted = group_name.strip().casefold()
        if not wanted:
            raise HTTPException(status_code=422, detail="Selecciona un grupo")
        selected = []
        for person in candidates:
            try:
                metadata = json.loads(person.metadata_json or "{}")
            except ValueError:
                metadata = {}
            if isinstance(metadata, dict) and str(metadata.get("group_name") or "").strip().casefold() == wanted:
                selected.append(person)
    elif mode == "excel":
        if upload is None:
            raise HTTPException(status_code=422, detail="Selecciona el archivo Excel")
        field, keys = _excel_assignment_keys(upload.filename or "", await upload.read())
        by_key: dict[str, list[Participant]] = {}
        for person in candidates:
            key = str(getattr(person, field) or "").strip().casefold()
            if key:
                by_key.setdefault(key, []).append(person)
        selected = []
        seen: set[str] = set()
        for row_number, key in enumerate(keys, 2):
            normalized = key.casefold()
            if normalized in seen:
                issues.append(f"Fila {row_number}: identificador repetido: {key}")
                continue
            seen.add(normalized)
            matches = by_key.get(normalized, [])
            if len(matches) != 1:
                issues.append(f"Fila {row_number}: {key} {'no pertenece a la fuente del formulario' if not matches else 'coincide con varios participantes'}")
            else:
                selected.append(matches[0])
    else:
        raise HTTPException(status_code=422, detail="Elige grupo o Excel")
    if not selected:
        issues.append("No hay participantes para asignar")
    existing = {row.participant_id: row for row in db.query(ParticipantFormAssignment).filter(
        ParticipantFormAssignment.template_id == template_id,
        ParticipantFormAssignment.participant_id.in_([person.id for person in selected]),
    ).all()}
    record_ids = {row[0] for row in db.query(RuntimeRecord.participant_id).filter(
        RuntimeRecord.template_id == template_id,
        RuntimeRecord.participant_id.in_([person.id for person in selected]),
        RuntimeRecord.status.notin_(["voided", "cancelled"]),
    ).all()}
    result = BulkAssignmentPreview(selected=len(selected), to_assign=0, already_assigned=0, to_reassign=0, protected=0, issues=issues)
    actionable = []
    for person in selected:
        if not participant_visible(db, responsible_user_id, person):
            result.issues.append(f"{person.full_name}: fuera del territorio del responsable")
            continue
        previous = existing.get(person.id)
        if previous and (previous.status in {"completed", "closed"} or person.id in record_ids):
            result.protected += 1
        elif previous and previous.responsible_user_id == responsible_user_id:
            result.already_assigned += 1
        elif previous:
            result.to_reassign += 1
            actionable.append(person)
        else:
            result.to_assign += 1
            actionable.append(person)
    if preview_only or result.issues:
        return result
    for person in actionable:
        row = upsert_assignment(db, template, person.id, responsible_user_id, user.id)
        db.add(AuditLog(project_id=template.project_id, user_id=user.id, module="participants", action="bulk_assign_form",
                        entity_type="participant_form_assignment", entity_id=row.id,
                        after_json=json.dumps({"template_id": template_id, "participant_id": person.id,
                                               "responsible_user_id": responsible_user_id, "mode": mode})))
    db.commit()
    result.applied = len(actionable)
    return result


@router.post("/templates/{template_id}", response_model=AssignmentRead)
def assign_participant(template_id: str, payload: AssignmentWrite, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> AssignmentRead:
    template = _template(db, user, template_id)
    require_any_project_permission(db, user.id, template.project_id, {IDENTITY_USERS_MANAGE})
    if configured_source(template).access_mode == "legacy":
        raise HTTPException(status_code=422, detail="Selecciona primero el modo abierto o cerrado del formulario")
    if payload.participant_id not in {person.id for person in eligible_participants(db, template, enforce_assignments=False)}:
        raise HTTPException(status_code=422, detail="El participante no cumple la fuente configurada")
    participant = db.get(Participant, payload.participant_id)
    if not participant_visible(db, payload.responsible_user_id, participant):
        raise HTTPException(status_code=422, detail="El responsable no puede acceder al territorio del participante")
    row = upsert_assignment(db, template, payload.participant_id, payload.responsible_user_id, user.id)
    db.add(AuditLog(project_id=template.project_id, user_id=user.id, module="participants", action="assign_form",
                    entity_type="participant_form_assignment", entity_id=row.id,
                    after_json=json.dumps({"template_id": template_id, "participant_id": payload.participant_id,
                                           "responsible_user_id": payload.responsible_user_id})))
    db.commit()
    db.refresh(row)
    return _read(db, row)


@router.delete("/templates/{template_id}/{participant_id}")
def remove_assignment(template_id: str, participant_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict[str, bool]:
    template = _template(db, user, template_id)
    require_any_project_permission(db, user.id, template.project_id, {IDENTITY_USERS_MANAGE})
    row = db.query(ParticipantFormAssignment).filter_by(template_id=template_id, participant_id=participant_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Asignación no encontrada")
    if db.query(RuntimeRecord.id).filter_by(template_id=template_id, participant_id=participant_id).first():
        raise HTTPException(status_code=409, detail="La actividad tiene respuestas; reasígnala o ciérrala en lugar de eliminarla")
    db.delete(row)
    db.commit()
    return {"deleted": True}


@router.post("/templates/{template_id}/{participant_id}/start", response_model=AssignmentRead)
def start_assignment(template_id: str, participant_id: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> AssignmentRead:
    template = _template(db, user, template_id)
    require_any_project_permission(db, user.id, template.project_id, {RECORDS_WRITE})
    row = visible_assignment(db, template, participant_id, user.id)
    if row.responsible_user_id != user.id:
        raise HTTPException(status_code=403, detail="Actividad no disponible para este responsable")
    if db.query(RuntimeRecord.id).filter(RuntimeRecord.template_id == template_id,
        RuntimeRecord.participant_id == participant_id,
        RuntimeRecord.status.notin_(["voided", "cancelled"])).first():
        raise HTTPException(status_code=409, detail="Este participante ya tiene un registro asociado a esta actividad. Abre la respuesta para consultarla.")
    if row.status in {"completed", "closed"}:
        raise HTTPException(status_code=409, detail="Esta actividad ya está completada o cerrada para este participante.")
    if row.status == "assigned":
        row.status = "in_progress"
        row.updated_at = utc_now()
        db.commit()
    return _read(db, row)


@router.patch("/templates/{template_id}/{participant_id}/phone", response_model=ParticipantRead)
def update_assigned_participant_phone(template_id: str, participant_id: str, payload: ParticipantPhoneUpdate,
                                      db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> ParticipantRead:
    template = _template(db, user, template_id)
    require_any_project_permission(db, user.id, template.project_id, {RECORDS_WRITE})
    row = visible_assignment(db, template, participant_id, user.id)
    if row.responsible_user_id != user.id or row.status not in {"assigned", "in_progress"}:
        raise HTTPException(status_code=403, detail="El participante no está disponible para esta captura")
    if db.query(RuntimeRecord.id).filter(RuntimeRecord.template_id == template_id,
        RuntimeRecord.participant_id == participant_id,
        RuntimeRecord.status.notin_(["voided", "cancelled"])).first():
        raise HTTPException(status_code=409, detail="La respuesta ya fue enviada; solicita una corrección autorizada")
    participant = db.query(Participant).filter_by(id=participant_id, project_id=template.project_id).with_for_update().first()
    if not participant or not participant_visible(db, user.id, participant):
        raise HTTPException(status_code=404, detail="Participante no encontrado")
    try:
        metadata = json.loads(participant.metadata_json or "{}")
    except ValueError:
        metadata = {}
    if not isinstance(metadata, dict):
        metadata = {}
    phone_keys = ("phone", "telefono", "teléfono", "celular", "mobile", "phone_number")
    phone_key = next((key for key in phone_keys if key in metadata), "phone")
    previous = str(metadata.get(phone_key) or "")
    updated = payload.phone.strip()
    if previous != updated:
        metadata[phone_key] = updated
        participant.metadata_json = json.dumps(metadata, ensure_ascii=False)
        db.add(AuditLog(project_id=template.project_id, user_id=user.id, module="participants", action="update_phone_from_form",
                        entity_type="participant", entity_id=participant.id,
                        before_json=json.dumps({"phone": previous}, ensure_ascii=False),
                        after_json=json.dumps({"phone": updated, "template_id": template.id}, ensure_ascii=False)))
        db.commit()
        db.refresh(participant)
    return participant_to_read(participant)


@router.post("/templates/{template_id}/participants", response_model=AssignmentRead)
def create_participant_in_open_form(template_id: str, payload: OpenParticipantCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> AssignmentRead:
    template = _template(db, user, template_id)
    require_any_project_permission(db, user.id, template.project_id, {RECORDS_WRITE})
    require_any_project_permission(db, user.id, template.project_id, {PARTICIPANTS_CREATE, IDENTITY_USERS_MANAGE})
    source = configured_source(template)
    if source.access_mode != "open":
        raise HTTPException(status_code=403, detail="Este formulario no permite participantes nuevos")
    key = payload.document_id if source.participant_key_field == "document_id" else payload.external_code
    if not key or not key.strip():
        raise HTTPException(status_code=422, detail="Ingresa la llave configurada para este formulario")
    existing = db.query(Participant).filter(Participant.project_id == template.project_id,
        func.lower(getattr(Participant, source.participant_key_field)) == key.strip().lower()).first()
    if existing:
        raise HTTPException(status_code=409, detail="Ya existe un participante con esa llave")
    metadata = {"department": (payload.department or "").strip(), "municipality": (payload.municipality or "").strip()}
    participant = Participant(project_id=template.project_id, full_name=payload.full_name.strip(),
        document_id=(payload.document_id or "").strip() or None,
        external_code=(payload.external_code or "").strip() or None,
        metadata_json=json.dumps(metadata, ensure_ascii=False))
    if not participant_visible(db, user.id, participant):
        raise HTTPException(status_code=403, detail="Fuera del territorio asignado")
    db.add(participant)
    db.flush()
    row = ParticipantFormAssignment(project_id=template.project_id, template_id=template.id,
        participant_id=participant.id, responsible_user_id=user.id, status="assigned", assigned_by=user.id)
    db.add(row)
    db.add(AuditLog(project_id=template.project_id, user_id=user.id, module="participants", action="create_in_form",
                    entity_type="participant", entity_id=participant.id,
                    after_json=json.dumps({"template_id": template_id, "participant_id": participant.id})))
    db.commit()
    db.refresh(row)
    return _read(db, row)
