"""Gestor teams are organizational membership, not form access grants."""

import csv
import json
from io import BytesIO, StringIO
from uuid import uuid4

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.api.permissions import require_any_project_permission
from app.core.permissions import BUILDER_WRITE, IDENTITY_USERS_MANAGE
from app.db.session import get_db
from app.models.audit import AuditLog
from app.models.gestor_team import GestorTeam, GestorTeamParticipant, GestorTeamUser
from app.models.assignment import UserOrganizationAssignment, UserProjectAssignment
from app.models.identity import Project
from app.models.identity import User
from app.models.participants import Participant
from app.services.assignment_service import assignment_service
from openpyxl import load_workbook

router = APIRouter()


class TeamWrite(BaseModel):
    name: str = Field(min_length=2, max_length=120)


class TeamMembersWrite(BaseModel):
    user_ids: list[str] = Field(default_factory=list)
    participant_ids: list[str] = Field(default_factory=list)


class TeamMemberWrite(BaseModel):
    member_id: str


def _project_user_ids(db: Session, project_id: str) -> set[str]:
    ids = {row[0] for row in db.query(UserProjectAssignment.user_id).join(User, User.id == UserProjectAssignment.user_id).filter(
        UserProjectAssignment.project_id == project_id, UserProjectAssignment.status == "active", User.status == "active").all()}
    project = db.get(Project, project_id)
    if project and project.organization_id:
        ids.update(row[0] for row in db.query(UserOrganizationAssignment.user_id).join(User, User.id == UserOrganizationAssignment.user_id).filter(
            UserOrganizationAssignment.organization_id == project.organization_id,
            UserOrganizationAssignment.status == "active", User.status == "active").all())
    return ids


def _manage(db: Session, actor: User, project_id: str) -> None:
    if not assignment_service.user_has_project_access(db, actor.id, project_id):
        raise HTTPException(status_code=403, detail="Sin acceso al proyecto")
    require_any_project_permission(db, actor.id, project_id, {IDENTITY_USERS_MANAGE})


def _view(db: Session, actor: User, project_id: str) -> None:
    if not assignment_service.user_has_project_access(db, actor.id, project_id):
        raise HTTPException(status_code=403, detail="Sin acceso al proyecto")
    require_any_project_permission(db, actor.id, project_id, {IDENTITY_USERS_MANAGE, BUILDER_WRITE})


def _team(db: Session, actor: User, team_id: str) -> GestorTeam:
    team = db.get(GestorTeam, team_id)
    if not team:
        raise HTTPException(status_code=404, detail="Equipo no encontrado")
    _manage(db, actor, team.project_id)
    return team


def _read(db: Session, team: GestorTeam) -> dict:
    users = db.query(GestorTeamUser.user_id).filter_by(team_id=team.id).all()
    participants = db.query(GestorTeamParticipant.participant_id).filter_by(team_id=team.id).all()
    return {"id": team.id, "project_id": team.project_id, "name": team.name,
            "user_ids": [row[0] for row in users], "participant_ids": [row[0] for row in participants]}


@router.get("/project/{project_id}")
def list_teams(project_id: str, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> list[dict]:
    _manage(db, actor, project_id)
    return [_read(db, team) for team in db.query(GestorTeam).filter_by(project_id=project_id).order_by(GestorTeam.name).all()]


@router.get("/project/{project_id}/summaries")
def team_summaries(project_id: str, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> list[dict]:
    _view(db, actor, project_id)
    rows = db.query(GestorTeam).filter_by(project_id=project_id).order_by(GestorTeam.name).all()
    user_counts = dict(db.query(GestorTeamUser.team_id, func.count()).join(GestorTeam, GestorTeam.id == GestorTeamUser.team_id).filter(
        GestorTeam.project_id == project_id).group_by(GestorTeamUser.team_id).all())
    participant_counts = dict(db.query(GestorTeamParticipant.team_id, func.count()).join(GestorTeam, GestorTeam.id == GestorTeamParticipant.team_id).filter(
        GestorTeam.project_id == project_id).group_by(GestorTeamParticipant.team_id).all())
    return [{"id": row.id, "project_id": project_id, "name": row.name,
             "user_count": user_counts.get(row.id, 0), "participant_count": participant_counts.get(row.id, 0)} for row in rows]


@router.get("/{team_id}/members/search")
def search_members(team_id: str, kind: str = Query(..., pattern="^(gestor|participante)$"), q: str = "",
                   limit: int = Query(50, ge=1, le=100), db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> list[dict]:
    team = _team(db, actor, team_id)
    query = q.strip()
    if kind == "gestor":
        valid_ids = _project_user_ids(db, team.project_id)
        if not valid_ids:
            return []
        users = db.query(User).filter(User.id.in_(valid_ids))
        if query:
            pattern = f"%{query}%"
            users = users.filter(or_(User.full_name.ilike(pattern), User.email.ilike(pattern), User.document_id.ilike(pattern)))
        rows = users.order_by(User.full_name).limit(limit).all()
        selected = {row[0] for row in db.query(GestorTeamUser.user_id).filter(GestorTeamUser.team_id == team.id,
            GestorTeamUser.user_id.in_([person.id for person in rows])).all()}
        return [{"id": person.id, "name": person.full_name, "identifier": person.email, "selected": person.id in selected} for person in rows]
    participants = db.query(Participant).filter(Participant.project_id == team.project_id, Participant.status == "active")
    if query:
        pattern = f"%{query}%"
        participants = participants.filter(or_(Participant.full_name.ilike(pattern), Participant.document_id.ilike(pattern), Participant.external_code.ilike(pattern)))
    rows = participants.order_by(Participant.full_name).limit(limit).all()
    selected = {row[0] for row in db.query(GestorTeamParticipant.participant_id).filter(GestorTeamParticipant.team_id == team.id,
        GestorTeamParticipant.participant_id.in_([person.id for person in rows])).all()}
    return [{"id": person.id, "name": person.full_name, "identifier": person.document_id or person.external_code or "",
             "selected": person.id in selected} for person in rows]


@router.post("/{team_id}/members/{kind}")
def add_member(team_id: str, kind: str, payload: TeamMemberWrite, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> dict:
    team = _team(db, actor, team_id)
    if kind == "gestor":
        if payload.member_id not in _project_user_ids(db, team.project_id):
            raise HTTPException(status_code=422, detail="El gestor no está activo en el proyecto")
        model, column = GestorTeamUser, GestorTeamUser.user_id
        row = GestorTeamUser(team_id=team.id, user_id=payload.member_id)
    elif kind == "participante":
        person = db.get(Participant, payload.member_id)
        if not person or person.project_id != team.project_id or person.status != "active":
            raise HTTPException(status_code=422, detail="El participante no está activo en el proyecto")
        model, column = GestorTeamParticipant, GestorTeamParticipant.participant_id
        row = GestorTeamParticipant(team_id=team.id, participant_id=payload.member_id)
    else:
        raise HTTPException(status_code=422, detail="Tipo de integrante inválido")
    if not db.query(model.id).filter(model.team_id == team.id, column == payload.member_id).first():
        db.add(row)
        db.add(AuditLog(project_id=team.project_id, user_id=actor.id, module="participants", action="add_gestor_team_member",
                        entity_type="gestor_team", entity_id=team.id,
                        after_json=json.dumps({"kind": kind, "member_id": payload.member_id})))
        db.commit()
    return {"selected": True}


@router.delete("/{team_id}/members/{kind}/{member_id}")
def remove_member(team_id: str, kind: str, member_id: str, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> dict:
    team = _team(db, actor, team_id)
    if kind == "gestor":
        db.query(GestorTeamUser).filter_by(team_id=team.id, user_id=member_id).delete(synchronize_session=False)
    elif kind == "participante":
        db.query(GestorTeamParticipant).filter_by(team_id=team.id, participant_id=member_id).delete(synchronize_session=False)
    else:
        raise HTTPException(status_code=422, detail="Tipo de integrante inválido")
    db.add(AuditLog(project_id=team.project_id, user_id=actor.id, module="participants", action="remove_gestor_team_member",
                    entity_type="gestor_team", entity_id=team.id,
                    after_json=json.dumps({"kind": kind, "member_id": member_id})))
    db.commit()
    return {"selected": False}


def _bulk_rows(filename: str, content: bytes) -> list[tuple[int, str, str, str]]:
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="El archivo supera 10 MB")
    try:
        if filename.lower().endswith(".xlsx"):
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            rows = workbook.active.iter_rows(values_only=True)
        elif filename.lower().endswith(".csv"):
            rows = csv.reader(StringIO(content.decode("utf-8-sig")))
        else:
            raise HTTPException(status_code=422, detail="Sube un archivo .xlsx o .csv")
        headers = [str(value or "").strip().casefold() for value in next(rows, [])]
        if headers != ["equipo", "tipo", "identificador"]:
            raise HTTPException(status_code=422, detail="La plantilla debe tener las columnas equipo, tipo, identificador")
        result = []
        for number, row in enumerate(rows, 2):
            if number > 50001:
                raise HTTPException(status_code=422, detail="Máximo 50.000 filas por archivo")
            if not any(value is not None and str(value).strip() for value in row):
                continue
            values = [str(row[index]).strip() if index < len(row) and row[index] is not None else "" for index in range(3)]
            result.append((number, values[0], values[1].casefold(), values[2]))
        if not result:
            raise HTTPException(status_code=422, detail="El archivo no contiene integrantes")
        return result
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"No fue posible leer el archivo: {exc}") from exc


@router.get("/project/{project_id}/bulk-template")
def bulk_template(project_id: str, db: Session = Depends(get_db), actor: User = Depends(get_current_user)):
    _manage(db, actor, project_id)
    content = "equipo,tipo,identificador\r\nEquipo Soacha,gestor,gestor@ejemplo.com\r\nEquipo Soacha,participante,123456789\r\n"
    return StreamingResponse(iter(["\ufeff" + content]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": "attachment; filename=plantilla_equipos_gestores.csv"})


@router.post("/project/{project_id}/bulk-import")
async def bulk_import(project_id: str, preview_only: bool = Form(True), upload: UploadFile = File(...),
                      db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> dict:
    _manage(db, actor, project_id)
    rows = _bulk_rows(upload.filename or "", await upload.read())
    teams = {team.name.casefold(): team for team in db.query(GestorTeam).filter_by(project_id=project_id).all()}
    valid_users = _project_user_ids(db, project_id)
    user_keys: dict[str, set[str]] = {}
    for uid, email, document in db.query(User.id, User.email, User.document_id).filter(User.id.in_(valid_users)).all():
        for key in (email, document):
            if key:
                user_keys.setdefault(key.strip().casefold(), set()).add(uid)
    participant_keys: dict[str, set[str]] = {}
    for pid, document, code in db.query(Participant.id, Participant.document_id, Participant.external_code).filter(
        Participant.project_id == project_id, Participant.status == "active").all():
        for key in (document, code):
            if key:
                participant_keys.setdefault(key.strip().casefold(), set()).add(pid)
    issues: list[str] = []
    selected: set[tuple[str, str, str]] = set()
    new_teams: dict[str, str] = {}
    for number, team_name, kind, identifier in rows:
        if not team_name or len(team_name) > 120 or kind not in {"gestor", "participante"} or not identifier:
            issues.append(f"Fila {number}: equipo, tipo o identificador inválido")
            continue
        key = team_name.casefold()
        new_teams.setdefault(key, team_name)
        matches = (user_keys if kind == "gestor" else participant_keys).get(identifier.casefold(), set())
        if len(matches) != 1:
            issues.append(f"Fila {number}: identificador {identifier} {'no encontrado' if not matches else 'ambiguo'} para {kind}")
            continue
        selected.add((key, kind, next(iter(matches))))
    if len(issues) > 100:
        issues = issues[:100] + ["Se muestran solo los primeros 100 errores"]
    existing_users = {(team.name.casefold(), "gestor", uid) for team, uid in db.query(GestorTeam, GestorTeamUser.user_id).join(
        GestorTeamUser, GestorTeamUser.team_id == GestorTeam.id).filter(GestorTeam.project_id == project_id).all()}
    existing_participants = {(team.name.casefold(), "participante", pid) for team, pid in db.query(GestorTeam, GestorTeamParticipant.participant_id).join(
        GestorTeamParticipant, GestorTeamParticipant.team_id == GestorTeam.id).filter(GestorTeam.project_id == project_id).all()}
    to_add = selected - existing_users - existing_participants
    result = {"rows": len(rows), "teams_to_create": len(set(new_teams) - set(teams)),
              "new_members": len(to_add), "already_members": len(selected) - len(to_add),
              "issues": issues, "applied": 0}
    if preview_only or issues:
        return result
    try:
        for key in set(new_teams) - set(teams):
            team = GestorTeam(project_id=project_id, name=new_teams[key], created_by=actor.id)
            db.add(team)
            teams[key] = team
        db.flush()
        db.bulk_save_objects([GestorTeamUser(id=str(uuid4()), team_id=teams[key].id, user_id=uid)
                              for key, kind, uid in to_add if kind == "gestor"])
        db.bulk_save_objects([GestorTeamParticipant(id=str(uuid4()), team_id=teams[key].id, participant_id=pid)
                              for key, kind, pid in to_add if kind == "participante"])
        db.add(AuditLog(project_id=project_id, user_id=actor.id, module="participants", action="bulk_import_gestor_teams",
                        entity_type="gestor_team", entity_id=project_id,
                        after_json=json.dumps({"rows": len(rows), "teams_created": result["teams_to_create"],
                                               "members_added": len(to_add)})))
        db.commit()
    except Exception:
        db.rollback()
        raise
    result["applied"] = len(to_add)
    return result


@router.post("/project/{project_id}", status_code=201)
def create_team(project_id: str, payload: TeamWrite, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> dict:
    _manage(db, actor, project_id)
    name = payload.name.strip()
    if db.query(GestorTeam.id).filter(GestorTeam.project_id == project_id, func.lower(GestorTeam.name) == name.lower()).first():
        raise HTTPException(status_code=409, detail="Ya existe un equipo con ese nombre")
    team = GestorTeam(project_id=project_id, name=name, created_by=actor.id)
    db.add(team)
    db.flush()
    db.add(AuditLog(project_id=project_id, user_id=actor.id, module="participants", action="create_gestor_team",
                    entity_type="gestor_team", entity_id=team.id, after_json=json.dumps({"name": name})))
    db.commit()
    return _read(db, team)


@router.patch("/{team_id}")
def rename_team(team_id: str, payload: TeamWrite, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> dict:
    team = _team(db, actor, team_id)
    name = payload.name.strip()
    if db.query(GestorTeam.id).filter(GestorTeam.project_id == team.project_id, GestorTeam.id != team.id,
                                       func.lower(GestorTeam.name) == name.lower()).first():
        raise HTTPException(status_code=409, detail="Ya existe un equipo con ese nombre")
    before = team.name
    team.name = name
    db.add(AuditLog(project_id=team.project_id, user_id=actor.id, module="participants", action="rename_gestor_team",
                    entity_type="gestor_team", entity_id=team.id, before_json=json.dumps({"name": before}),
                    after_json=json.dumps({"name": name})))
    db.commit()
    return _read(db, team)


@router.put("/{team_id}/members")
def set_members(team_id: str, payload: TeamMembersWrite, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> dict:
    team = _team(db, actor, team_id)
    user_ids, participant_ids = set(payload.user_ids), set(payload.participant_ids)
    if len(user_ids) > 100 or len(participant_ids) > 10000:
        raise HTTPException(status_code=422, detail="Demasiados integrantes para un equipo")
    invalid_users = [uid for uid in user_ids if not db.get(User, uid) or not assignment_service.user_has_project_access(db, uid, team.project_id)]
    valid_participants = {row[0] for row in db.query(Participant.id).filter(Participant.project_id == team.project_id,
                                                                             Participant.status == "active", Participant.id.in_(participant_ids)).all()}
    if invalid_users or valid_participants != participant_ids:
        raise HTTPException(status_code=422, detail="Usuarios o participantes fuera del proyecto o inactivos")
    before = _read(db, team)
    db.query(GestorTeamUser).filter_by(team_id=team.id).delete(synchronize_session=False)
    db.query(GestorTeamParticipant).filter_by(team_id=team.id).delete(synchronize_session=False)
    db.add_all(GestorTeamUser(team_id=team.id, user_id=uid) for uid in user_ids)
    db.add_all(GestorTeamParticipant(team_id=team.id, participant_id=pid) for pid in participant_ids)
    db.flush()
    after = _read(db, team)
    db.add(AuditLog(project_id=team.project_id, user_id=actor.id, module="participants", action="set_gestor_team_members",
                    entity_type="gestor_team", entity_id=team.id, before_json=json.dumps(before), after_json=json.dumps(after)))
    db.commit()
    return after


@router.delete("/{team_id}")
def delete_team(team_id: str, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> dict[str, bool]:
    team = _team(db, actor, team_id)
    before = _read(db, team)
    db.query(GestorTeamUser).filter_by(team_id=team.id).delete(synchronize_session=False)
    db.query(GestorTeamParticipant).filter_by(team_id=team.id).delete(synchronize_session=False)
    db.add(AuditLog(project_id=team.project_id, user_id=actor.id, module="participants", action="delete_gestor_team",
                    entity_type="gestor_team", entity_id=team.id, before_json=json.dumps(before)))
    db.delete(team)
    db.commit()
    return {"deleted": True}
