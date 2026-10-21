"""Devolución y reasignación masiva desde Excel o CSV (docs/141 §5).

Columnas: id_registro, accion (devolver_original | reasignar), usuario_destino
(correo o documento; obligatorio al reasignar), observacion (concepto
general), campo (pregunta o ruta REPEAT, p. ej. integrantes[2].edad) y motivo.
Varias filas del mismo registro se aplican como UNA devolución con varias
observaciones.

Se valida todo antes de aplicar (existencia, proyecto, territorio, estado,
destinatario y duplicados). Aplicar es idempotente: cada fila tiene una
referencia estable (registro + campo + motivo), así que volver a cargar el
mismo archivo no duplica observaciones ni devoluciones. Un registro con
error no detiene a los demás.
"""

from __future__ import annotations

import csv
import hashlib
from io import BytesIO, StringIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.api.permissions import participant_visible
from app.models.identity import User
from app.models.participants import Participant
from app.models.review import ReviewObservation
from app.models.runtime_record import RuntimeRecord
from app.services.correction_service import return_record, validate_assignee
from app.services.review_service import ALLOWED_TRANSITIONS

MAX_ROWS = 5000
COLUMNS = ["id_registro", "accion", "usuario_destino", "observacion", "campo", "motivo"]
ALIASES = {
    "id registro": "id_registro", "registro": "id_registro", "record_id": "id_registro", "id": "id_registro",
    "acción": "accion", "action": "accion",
    "usuario destino": "usuario_destino", "destino": "usuario_destino", "destinatario": "usuario_destino",
    "observación": "observacion", "concepto": "observacion",
    "pregunta": "campo", "campo afectado": "campo", "field": "campo",
}
RETURN_TO_AUTHOR = {"devolver_original", "devolver", "original", "devolver al original"}
REASSIGN = {"reasignar", "reasignar_corrección", "reasignar_correccion"}


class BulkReturnError(ValueError):
    pass


def read_rows(filename: str, content: bytes) -> list[tuple[int, dict[str, str]]]:
    if len(content) > 10 * 1024 * 1024:
        raise BulkReturnError("El archivo supera 10 MB")
    name = filename.lower()
    try:
        if name.endswith(".xlsx"):
            rows = load_workbook(BytesIO(content), read_only=True, data_only=True).active.iter_rows(values_only=True)
        elif name.endswith(".csv"):
            rows = csv.reader(StringIO(content.decode("utf-8-sig")))
        else:
            raise BulkReturnError("Sube un archivo .xlsx o .csv")
        headers = [ALIASES.get(str(v or "").strip().casefold(), str(v or "").strip().casefold().replace(" ", "_")) for v in next(rows, [])]
    except BulkReturnError:
        raise
    except Exception as exc:  # archivo dañado o formato desconocido
        raise BulkReturnError(f"No fue posible leer el archivo: {exc}") from exc
    if "id_registro" not in headers or "accion" not in headers:
        raise BulkReturnError("La plantilla debe tener al menos las columnas id_registro y accion")
    result: list[tuple[int, dict[str, str]]] = []
    for number, row in enumerate(rows, 2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        if len(result) >= MAX_ROWS:
            raise BulkReturnError(f"Máximo {MAX_ROWS} filas por archivo")
        result.append((number, {headers[i]: (str(row[i]).strip() if i < len(row) and row[i] is not None else "") for i in range(len(headers))}))
    if not result:
        raise BulkReturnError("El archivo no tiene filas")
    return result


def _row_ref(record_id: str, field: str, reason: str) -> str:
    return "bulk-" + hashlib.sha256(f"{record_id}|{field}|{reason}".encode("utf-8")).hexdigest()[:40]


def _resolve_user(db: Session, value: str) -> User | None:
    value = value.strip()
    if not value:
        return None
    return (db.query(User).filter(User.email == value.lower()).first()
            or db.query(User).filter(User.document_id == value).first()
            or db.get(User, value))


def process(db: Session, project_id: str, actor_id: str, rows: list[tuple[int, dict[str, str]]], *, preview_only: bool, apply_action) -> dict:
    results: dict[int, dict] = {}
    groups: dict[str, dict] = {}
    seen_refs: set[str] = set()

    for number, row in rows:
        record_id = row.get("id_registro", "")
        action = row.get("accion", "").strip().casefold()
        field = row.get("campo", "")
        reason = row.get("motivo", "") or row.get("observacion", "")
        out = {"fila": number, "id_registro": record_id, "estado": "ok", "detalle": ""}
        results[number] = out

        def fail(message: str) -> None:
            out.update(estado="error", detalle=message)

        record = db.get(RuntimeRecord, record_id) if record_id else None
        participant = db.get(Participant, record.participant_id) if record and record.participant_id else None
        if record is None or record.project_id != project_id or (participant is not None and not participant_visible(db, actor_id, participant)):
            fail("Registro no encontrado en el proyecto o fuera de su territorio")
            continue
        if action not in RETURN_TO_AUTHOR | REASSIGN:
            fail("Acción no válida: use devolver_original o reasignar")
            continue
        if not reason:
            fail("Falta el motivo u observación")
            continue
        assignee = None
        if action in REASSIGN:
            assignee = _resolve_user(db, row.get("usuario_destino", ""))
            if assignee is None:
                fail("Usuario de destino no encontrado")
                continue
            try:
                validate_assignee(db, record, assignee.id)
            except ValueError as exc:
                fail(str(exc))
                continue
        ref = _row_ref(record.id, field, reason)
        if ref in seen_refs:
            out.update(estado="duplicada", detalle="Fila repetida en el archivo; se omite")
            continue
        seen_refs.add(ref)
        group = groups.setdefault(record.id, {"record": record, "assignee": assignee.id if assignee else None, "action": action,
                                              "notes": row.get("observacion", "") or reason, "rows": [], "observations": []})
        if (group["assignee"], group["action"] in REASSIGN) != ((assignee.id if assignee else None), action in REASSIGN):
            fail("Otra fila de este registro pide una acción o destinatario distinto")
            continue
        group["rows"].append(number)
        group["observations"].append({"field_name": field.split(".")[-1].split("[")[0] if field else None,
                                      "repeat_path": field if "[" in field else None, "reason": reason, "client_ref": ref})

    # Estado del registro e idempotencia (se evalúa por registro)
    for record_id, group in list(groups.items()):
        record = group["record"]
        refs = [o["client_ref"] for o in group["observations"]]
        existing = {r for (r,) in db.query(ReviewObservation.client_ref).filter(ReviewObservation.record_id == record_id, ReviewObservation.client_ref.in_(refs)).all()}
        if record.status in {"returned", "corrected"} and set(refs) <= existing:
            for n in group["rows"]:
                results[n].update(estado="ya_aplicada", detalle="Ya se había devuelto con estas observaciones")
            groups.pop(record_id)
        elif "returned" not in ALLOWED_TRANSITIONS.get(record.status, set()):
            for n in group["rows"]:
                results[n].update(estado="error", detalle=f"El registro está en estado '{record.status}' y no se puede devolver")
            groups.pop(record_id)

    applied = 0
    if not preview_only:
        for record_id, group in groups.items():
            try:
                return_record(db, group["record"], actor_id, group["notes"], group["assignee"], group["observations"], apply_action)
                applied += 1
                for n in group["rows"]:
                    results[n].update(estado="aplicada", detalle="Devuelto" + (" y reasignado" if group["assignee"] else " al autor"))
            except ValueError as exc:
                db.rollback()
                for n in group["rows"]:
                    results[n].update(estado="error", detalle=str(exc))

    ordered = [results[n] for n, _ in rows]
    errors = [r for r in ordered if r["estado"] == "error"]
    errors_csv = StringIO()
    writer = csv.writer(errors_csv, lineterminator="\n")
    writer.writerow(["fila", "id_registro", "error"])
    for r in errors:
        writer.writerow([r["fila"], r["id_registro"], r["detalle"]])
    return {
        "filas": len(ordered),
        "registros_a_devolver": len(groups) if preview_only else applied,
        "aplicadas": applied,
        "errores": len(errors),
        "duplicadas": sum(1 for r in ordered if r["estado"] == "duplicada"),
        "ya_aplicadas": sum(1 for r in ordered if r["estado"] == "ya_aplicada"),
        "resultados": ordered,
        "errores_csv": errors_csv.getvalue() if errors else "",
        "vista_previa": preview_only,
    }


def template_csv() -> str:
    output = StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(COLUMNS)
    writer.writerow(["<id del registro>", "devolver_original", "", "Fecha y firma incorrectas", "fecha_visita", "La fecha no existe"])
    writer.writerow(["<id del registro>", "reasignar", "gestora@ejemplo.org", "Corregir integrante", "integrantes[2].edad", "Edad fuera de rango"])
    return "﻿" + output.getvalue()
