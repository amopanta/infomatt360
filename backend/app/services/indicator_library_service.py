"""Reusable indicator definitions and spreadsheet import."""

from io import BytesIO
import re
import unicodedata

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.time import utc_now
from app.models.builder import BuilderComponent, BuilderTemplate
from app.models.reports import ReportIndicator
from app.schemas.report_catalog import CatalogReportConfig, IndicatorDefinition, IndicatorSource, LibraryIndicatorRead
from app.services.report_catalog_service import report_catalog_service


def _normal(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip().casefold())
    return "".join(char for char in text if not unicodedata.combining(char))


class IndicatorLibraryService:
    def read(self, row: ReportIndicator) -> LibraryIndicatorRead:
        return LibraryIndicatorRead(id=row.id, project_id=row.project_id, definition=IndicatorDefinition.model_validate_json(row.definition_json), created_at=row.created_at, updated_at=row.updated_at)

    def validate(self, db: Session, project_id: str, definition: IndicatorDefinition) -> None:
        if not definition.code.strip():
            raise HTTPException(status_code=422, detail="El código del indicador es obligatorio")
        report_catalog_service.validate(db, CatalogReportConfig(project_id=project_id, name="Validación", indicators=[definition]))

    def save(self, db: Session, project_id: str, definition: IndicatorDefinition, indicator_id: str | None = None) -> LibraryIndicatorRead:
        self.validate(db, project_id, definition)
        row = db.query(ReportIndicator).filter(ReportIndicator.id == indicator_id, ReportIndicator.project_id == project_id).first() if indicator_id else None
        if indicator_id and row is None:
            raise HTTPException(status_code=404, detail="Indicador no encontrado")
        duplicate = db.query(ReportIndicator).filter(ReportIndicator.project_id == project_id, ReportIndicator.code == definition.code.strip()).first()
        if duplicate and (row is None or duplicate.id != row.id):
            raise HTTPException(status_code=409, detail=f"Ya existe el indicador {definition.code}")
        if row is None:
            row = ReportIndicator(project_id=project_id, code=definition.code.strip(), title=definition.title, definition_json=definition.model_dump_json())
            db.add(row)
        else:
            row.code = definition.code.strip()
            row.title = definition.title
            row.definition_json = definition.model_dump_json()
            row.updated_at = utc_now()
        db.commit()
        db.refresh(row)
        return self.read(row)

    def template(self) -> bytes:
        book = Workbook()
        sheet = book.active
        sheet.title = "indicadores"
        sheet.append(["Código", "Indicador", "Meta", "Unidad", "Formularios asociados", "Campo de cálculo", "Operación", "Estado requerido", "Avance manual"])
        sheet.append(["IND-001", "Familias con visita", 4000, "Familias", "Nombre del formulario", "codigo_participante", "union", "approved", None])
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(18, len(str(column[0].value)) + 3))
        output = BytesIO()
        book.save(output)
        return output.getvalue()

    def parse(self, db: Session, project_id: str, content: bytes) -> tuple[list[IndicatorDefinition], list[dict]]:
        try:
            book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"No se pudo leer el Excel: {exc}") from exc
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headers = [_normal(value) for value in next(rows, [])]
        positions = {name: index for index, name in enumerate(headers)}
        if not {"codigo", "indicador", "meta"}.issubset(positions):
            raise HTTPException(status_code=422, detail="El Excel necesita Código, Indicador y Meta")
        templates = db.query(BuilderTemplate).filter(BuilderTemplate.project_id == project_id).all()
        by_name = {_normal(row.name): row for row in templates}
        by_id = {row.id: row for row in templates}
        definitions: list[IndicatorDefinition] = []
        errors: list[dict] = []
        seen_codes: set[str] = set()
        for row_number, row in enumerate(rows, 2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            def cell(name: str):
                index = positions.get(name)
                return row[index] if index is not None and index < len(row) else None
            try:
                code = str(cell("codigo") or "").strip()
                title = str(cell("indicador") or "").strip()
                goal = float(cell("meta"))
                if not code or not title or code.casefold() in seen_codes:
                    raise ValueError("Código o nombre vacío, o código repetido")
                seen_codes.add(code.casefold())
                forms = [part.strip() for part in re.split(r"[+;,]", str(cell("formularios asociados") or "")) if part.strip()]
                key_name = str(cell("campo de calculo") or "").strip()
                operation = _normal(cell("operacion") or "union")
                operation = {"conteo unico": "union", "unico": "union", "interseccion": "intersection", "cumplimiento de todos": "all", "suma": "sum"}.get(operation, operation)
                if operation not in {"sum", "union", "intersection", "all"}:
                    raise ValueError(f"Operación desconocida: {operation}")
                if forms:
                    sources = []
                    for form_name in forms:
                        template = by_id.get(form_name) or by_name.get(_normal(form_name))
                        if template is None:
                            raise ValueError(f"Formulario no encontrado: {form_name}")
                        field_name = None
                        if key_name:
                            field = db.query(BuilderComponent).filter(BuilderComponent.template_id == template.id).all()
                            matched = [component for component in field if _normal(component.name) == _normal(key_name) or _normal(component.label) == _normal(key_name)]
                            if len(matched) != 1:
                                raise ValueError(f"Llave {key_name} no encontrada o ambigua en {template.name}")
                            field_name = matched[0].name
                        sources.append(IndicatorSource(template_id=template.id, key_field=field_name, required_status=str(cell("estado requerido") or "").strip() or None))
                    definition = IndicatorDefinition(code=code, title=title, goal=goal, unit=str(cell("unidad") or ""), sources=sources, combination=operation)
                else:
                    definition = IndicatorDefinition(code=code, title=title, goal=goal, unit=str(cell("unidad") or ""), source_mode="manual", manual_actual=float(cell("avance manual") or 0))
                self.validate(db, project_id, definition)
                definitions.append(definition)
            except (ValueError, TypeError, HTTPException) as exc:
                errors.append({"row": row_number, "error": str(exc.detail) if isinstance(exc, HTTPException) else str(exc)})
        return definitions, errors

    def import_rows(self, db: Session, project_id: str, content: bytes) -> list[LibraryIndicatorRead]:
        definitions, errors = self.parse(db, project_id, content)
        if errors or not definitions:
            raise HTTPException(status_code=422, detail={"errors": errors or [{"row": 0, "error": "Sin indicadores"}]})
        codes = [item.code for item in definitions]
        existing = db.query(ReportIndicator.code).filter(ReportIndicator.project_id == project_id, ReportIndicator.code.in_(codes)).all()
        if existing:
            raise HTTPException(status_code=409, detail=f"Ya existen códigos: {', '.join(code for (code,) in existing)}")
        rows = [ReportIndicator(project_id=project_id, code=item.code, title=item.title, definition_json=item.model_dump_json()) for item in definitions]
        db.add_all(rows)
        db.commit()
        for row in rows:
            db.refresh(row)
        return [self.read(row) for row in rows]


indicator_library_service = IndicatorLibraryService()
