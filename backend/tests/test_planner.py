"""Planeador mensual (docs/144): criterios de aceptación del diagnóstico §7 y §17."""

import io
import json
from datetime import datetime

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.testclient import TestClient

from app.core.permissions import ALL_PERMISSIONS, IDENTITY_USERS_MANAGE, ORGANIZATIONS_MANAGE, PLANNER_MANAGE
from app.core.security import hash_password
from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.assignment import UserProjectAssignment
from app.models.builder import BuilderTemplate
from app.models.case_management import UserTerritory
from app.models.gestor_team import GestorTeam, GestorTeamUser
from app.models.identity import Project, Role, User
from app.models.participants import Participant
from app.models.planner import PlannerHistory, PlannerTarget
from app.models.runtime_record import RuntimeRecord, RuntimeRecordValue

P = "proj-plan"
PASSWORD = "Clave12345!"
NOV = datetime(2026, 11, 10, 15, 0)


def _seed(db):
    meta = lambda mun: json.dumps({"department": "Antioquia", "municipality": mun})  # noqa: E731
    field = ALL_PERMISSIONS - {IDENTITY_USERS_MANAGE, ORGANIZATIONS_MANAGE, PLANNER_MANAGE}
    db.add_all([
        Project(id=P, name="Proyecto planeador"),
        Project(id="otro", name="Otro"),
        Role(id="r-admin", name="Admin", permissions=",".join(sorted(ALL_PERMISSIONS))),
        Role(id="r-coord", name="Coordinador", permissions=",".join(sorted(field))),
        Role(id="r-gestor", name="Gestor", permissions="projects.read,records.write,records.read"),
        User(id="admin", full_name="Admin", document_id="A", email="admin@x.org", password_hash=hash_password(PASSWORD)),
        User(id="coord", full_name="Coordinadora Turbo", document_id="C", email="coord@x.org", password_hash=hash_password(PASSWORD)),
        User(id="g1", full_name="Gestor Uno", document_id="G1", email="g1@x.org", password_hash=hash_password(PASSWORD)),
        User(id="g2", full_name="Gestora Dos", document_id="G2", email="g2@x.org", password_hash=hash_password(PASSWORD)),
        User(id="ajeno", full_name="Ajeno", document_id="Z", email="ajeno@x.org", password_hash=hash_password(PASSWORD)),
        UserProjectAssignment(user_id="admin", project_id=P, role_id="r-admin", status="active"),
        UserProjectAssignment(user_id="coord", project_id=P, role_id="r-coord", status="active"),
        UserProjectAssignment(user_id="g1", project_id=P, role_id="r-gestor", status="active"),
        UserProjectAssignment(user_id="g2", project_id=P, role_id="r-gestor", status="active"),
        UserProjectAssignment(user_id="ajeno", project_id="otro", role_id="r-admin", status="active"),
        UserTerritory(project_id=P, user_id="coord", department="Antioquia", municipality="Turbo"),
        UserTerritory(project_id=P, user_id="g1", department="Antioquia", municipality="Turbo"),
        UserTerritory(project_id=P, user_id="g2", department="Antioquia", municipality="Apartadó"),
        GestorTeam(id="team-1", project_id=P, name="Equipo Turbo", created_by="admin", coordinator_user_id="coord"),
        GestorTeamUser(team_id="team-1", user_id="g1"),
        BuilderTemplate(id="v3", project_id=P, name="Visita 3 FSC", status="published"),
        BuilderTemplate(id="v4", project_id=P, name="Visita 4 VMC", status="published"),
    ])
    # 50 familias en Turbo; Apartadó tiene 2.
    for i in range(50):
        db.add(Participant(id=f"t{i}", project_id=P, full_name=f"Familia T{i}", document_id=f"D{i}", metadata_json=meta("Turbo")))
    db.add_all([Participant(id=f"a{i}", project_id=P, full_name=f"Familia A{i}", document_id=f"DA{i}", metadata_json=meta("Apartadó")) for i in range(2)])
    n = 0

    def record(template, participant, status, by, when=NOV, **values):
        nonlocal n
        n += 1
        db.add(RuntimeRecord(id=f"r{n}", project_id=P, template_id=template, participant_id=participant, status=status,
                             submitted_by=by, created_at=when, updated_at=when))
        for name, value in values.items():
            db.add(RuntimeRecordValue(id=f"v{n}-{name}", record_id=f"r{n}", field_name=name, field_value_json=json.dumps(value)))

    # Visita 3 en Turbo (g1): 85 aprobadas sobre 50 familias (35 con dos visitas) y 15 enviadas sin aprobar.
    for i in range(85):
        record("v3", f"t{i % 50}", "approved", "g1")
    for i in range(15):
        record("v3", f"t{i}", "submitted", "g1")
    # Visita 4 en Turbo (g1): 80 aprobadas.
    for i in range(80):
        record("v4", f"t{i % 50}", "approved", "g1")
    # Apartadó (g2): 3 aprobadas, una con fecha de actividad en octubre aunque se registró en noviembre.
    record("v3", "a0", "approved", "g2")
    record("v3", "a1", "approved", "g2")
    record("v3", "a1", "approved", "g2", fecha_actividad="2026-10-30")
    # No cuentan: borrador, anulado, otro año.
    record("v3", "t1", "draft", "g1")
    record("v3", "t2", "voided", "g1")
    record("v3", "t3", "approved", "g1", when=datetime(2025, 11, 10))
    db.commit()


@pytest.fixture()
def client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    sessions = sessionmaker(bind=engine)
    Base.metadata.create_all(bind=engine)
    with sessions() as db:
        _seed(db)

    def override_db():
        with sessions() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    try:
        with TestClient(app) as test_client:
            test_client.sessions = sessions
            yield test_client
    finally:
        app.dependency_overrides.clear()
        Base.metadata.drop_all(bind=engine)


def login(client, email):
    token = client.post("/api/v1/auth/login", json={"email": email, "password": PASSWORD}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def put_target(client, headers, template, target, municipality="Turbo", month=11, unit="records", gestor=None, reason=None):
    return client.put(f"/api/v1/planner/{P}/targets", headers=headers, json={
        "year": 2026, "month": month, "municipality": municipality, "template_id": template, "unit": unit,
        "target": target, "gestor_user_id": gestor, "reason": reason})


def municipal(report, activity, municipality="Turbo", month=11):
    return next(r for r in report["municipal"] if r["activity"] == activity and r["municipality"] == municipality and r["month"] == month)


def dashboard(client, headers, **params):
    response = client.get(f"/api/v1/planner/{P}/dashboard", headers=headers, params={"year": 2026, **params})
    assert response.status_code == 200, response.text
    return response.json()


def test_compliance_weighted_totals_and_no_goal(client):
    admin = login(client, "admin@x.org")
    assert put_target(client, admin, "v3", 100).status_code == 200
    assert put_target(client, admin, "v4", 80).status_code == 200
    report = dashboard(client, admin, month=11, municipality="Turbo")
    v3, v4 = municipal(report, "Visita 3 FSC"), municipal(report, "Visita 4 VMC")
    assert v3["activities"] == {"target": 100, "executed": 85, "pending": 15, "over": 0, "compliance": 85.0}
    assert v4["activities"]["compliance"] == 100.0
    # Ponderado: 165 / 180, nunca el promedio de 85 % y 100 %.
    assert report["totals"]["activities"]["compliance"] == 91.67
    # Apartadó no tiene meta: «Sin meta», no 0 % ni 100 %.
    apartado = municipal(dashboard(client, admin, month=11), "Visita 3 FSC", municipality="Apartadó")
    # Sin campo de fecha configurado, la visita fechada en octubre cuenta por su fecha de registro.
    assert apartado["activities"]["compliance"] is None and apartado["activities"]["executed"] == 3


def test_count_policy_is_configurable(client):
    admin = login(client, "admin@x.org")
    put_target(client, admin, "v3", 100)
    settings = {"count_policy": "submitted", "family_key": "participant", "timezone": "America/Bogota"}
    assert client.put(f"/api/v1/planner/{P}/settings", headers=admin, json=settings).status_code == 200
    v3 = municipal(dashboard(client, admin, month=11), "Visita 3 FSC")
    assert v3["activities"]["executed"] == 100 and v3["activities"]["compliance"] == 100.0
    gestor = login(client, "g1@x.org")
    assert client.put(f"/api/v1/planner/{P}/settings", headers=gestor, json=settings).status_code == 403


def test_unique_families_versus_visits_and_family_key(client):
    admin = login(client, "admin@x.org")
    v3 = municipal(dashboard(client, admin, month=11), "Visita 3 FSC")
    assert v3["activities"]["executed"] == 85 and v3["families"]["executed"] == 50
    # La llave de familia también puede ser un campo del formulario.
    settings = {"count_policy": "approved", "family_key": "field", "family_field": "codigo_hogar"}
    assert client.put(f"/api/v1/planner/{P}/settings", headers=admin, json=settings).status_code == 200
    assert municipal(dashboard(client, admin, month=11), "Visita 3 FSC")["families"]["executed"] == 0
    assert client.put(f"/api/v1/planner/{P}/settings", headers=admin, json={"count_policy": "approved", "family_key": "field"}).status_code == 422


def test_activity_date_field_and_population_coverage(client):
    admin = login(client, "admin@x.org")
    settings = {"count_policy": "approved", "family_key": "participant", "activity_date_field": "fecha_actividad"}
    client.put(f"/api/v1/planner/{P}/settings", headers=admin, json=settings)
    october = dashboard(client, admin, month=10)
    assert municipal(october, "Visita 3 FSC", municipality="Apartadó", month=10)["activities"]["executed"] == 1
    assert client.put(f"/api/v1/planner/{P}/population", headers=admin,
                      json={"year": 2026, "municipality": "Turbo", "families": 100, "participants": 400}).status_code == 200
    report = dashboard(client, admin, municipality="Turbo")
    assert report["totals"]["families"]["population"] == 100 and report["totals"]["families"]["coverage"] == 50.0
    # Cambiar la población exige motivo y queda en el historial.
    assert client.put(f"/api/v1/planner/{P}/population", headers=admin,
                      json={"year": 2026, "municipality": "Turbo", "families": 120}).status_code == 422


CSV_OK = ("anio,mes,municipio,formulario,gestor,unidad,meta,modo,justificacion\n"
          "2026,11,Turbo,Visita 3 FSC,,registros,100,crear,\n"
          "2026,11,Turbo,Visita 4 VMC,g1@x.org,registros,80,crear,\n"
          "2026,11,Turbo,Visita 3 FSC,,familias,50,crear,\n")


def upload(client, headers, text, preview=True, name="metas.csv"):
    return client.post(f"/api/v1/planner/{P}/targets/import", headers=headers, data={"preview_only": str(preview).lower()},
                       files={"upload": (name, text.encode("utf-8"), "text/csv")})


def test_import_preview_apply_and_repeat_does_not_duplicate(client):
    admin = login(client, "admin@x.org")
    preview = upload(client, admin, CSV_OK).json()
    assert [r["status"] for r in preview["rows"]] == ["create", "create", "create"] and preview["total_after"] == 230
    assert upload(client, admin, CSV_OK, preview=False).json()["applied"] == 3
    again = upload(client, admin, CSV_OK, preview=False).json()
    assert again["applied"] == 0 and {r["status"] for r in again["rows"]} == {"already_applied"}
    adding = CSV_OK.replace("100,crear,", "20,sumar,Refuerzo")
    first = upload(client, admin, adding, preview=False).json()
    assert first["applied"] == 1
    assert upload(client, admin, adding, preview=False).json()["applied"] == 0  # misma carga: no suma dos veces
    with client.sessions() as db:
        rows = db.query(PlannerTarget).filter(PlannerTarget.project_id == P).all()
        assert len(rows) == 3 and sorted(r.target for r in rows) == [50, 80, 120]
        history = db.query(PlannerHistory).filter(PlannerHistory.entity_type == "target").count()
        assert history == 4  # 3 creadas + 1 suma


def test_import_validations(client):
    admin = login(client, "admin@x.org")
    assert upload(client, admin, "anio,mes,municipio\n2026,11,Turbo\n").status_code == 422  # faltan columnas
    bad = ("anio,mes,municipio,formulario,gestor,unidad,meta,modo,justificacion\n"
           "2026,13,Turbo,Visita 3 FSC,,registros,10,crear,\n"
           "2026,11,Turbo,No existe,,registros,10,crear,\n"
           "2026,11,Turbo,Visita 3 FSC,nadie@x.org,registros,10,crear,\n"
           "2026,11,Turbo,Visita 3 FSC,,registros,-1,crear,\n"
           "2026,11,Turbo,Visita 3 FSC,,registros,10,crear,\n"
           "2026,11,Turbo,Visita 3 FSC,,registros,12,crear,\n")
    result = upload(client, admin, bad).json()
    assert [r["status"] for r in result["rows"]] == ["error"] * 4 + ["create", "error"]
    assert "Duplicada" in result["rows"][5]["message"]
    assert upload(client, admin, bad, preview=False).status_code == 422  # con errores no se aplica nada
    upload(client, admin, CSV_OK, preview=False)
    replace = CSV_OK.splitlines()[0] + "\n2026,11,Turbo,Visita 3 FSC,,registros,90,reemplazar,\n"
    assert upload(client, admin, replace).json()["rows"][0]["status"] == "error"  # reemplazar exige justificación
    # La coordinadora de Turbo no gestiona metas (no tiene planner.manage).
    assert upload(client, login(client, "coord@x.org"), CSV_OK).status_code == 403


def test_import_rejects_unauthorized_municipality(client):
    with client.sessions() as db:
        role = db.get(Role, "r-coord")
        role.permissions += f",{PLANNER_MANAGE}"
        db.commit()
    coord = login(client, "coord@x.org")
    text = CSV_OK.replace("Visita 4 VMC,g1@x.org", "Visita 4 VMC,").replace("Turbo,Visita 4", "Apartadó,Visita 4")
    rows = upload(client, coord, text).json()["rows"]
    assert rows[1]["status"] == "error" and "territorio" in rows[1]["message"]


def test_closed_cut_is_frozen(client):
    admin = login(client, "admin@x.org")
    put_target(client, admin, "v3", 100)
    before = municipal(dashboard(client, admin, month=11), "Visita 3 FSC")
    assert client.post(f"/api/v1/planner/{P}/cuts", headers=admin, json={"year": 2026, "month": 11}).status_code == 201
    assert client.post(f"/api/v1/planner/{P}/cuts", headers=admin, json={"year": 2026, "month": 11}).status_code == 409
    assert put_target(client, admin, "v3", 150, reason="Ajuste").status_code == 409
    with client.sessions() as db:  # aunque llegue un registro nuevo, el corte no cambia
        db.add(RuntimeRecord(id="tarde", project_id=P, template_id="v3", participant_id="t1", status="approved", submitted_by="g1",
                             created_at=NOV, updated_at=NOV))
        db.commit()
    after = municipal(dashboard(client, admin, month=11), "Visita 3 FSC")
    assert after["activities"] == before["activities"] and after["closed"] is True


def test_profiles_see_only_their_scope(client):
    admin = login(client, "admin@x.org")
    put_target(client, admin, "v3", 100)
    put_target(client, admin, "v3", 5, municipality="Apartadó")
    coord = dashboard(client, login(client, "coord@x.org"), month=11)
    assert {r["municipality"] for r in coord["municipal"]} == {"Turbo"}
    g2 = dashboard(client, login(client, "g2@x.org"), month=11)
    assert {r["gestor"] for r in g2["gestor"]} == {"Gestora Dos"}
    assert all(r["municipality"] == "Apartadó" for r in g2["municipal"])
    ajeno = login(client, "ajeno@x.org")
    assert client.get(f"/api/v1/planner/{P}/dashboard", headers=ajeno, params={"year": 2026}).status_code == 403
    assert client.get(f"/api/v1/planner/{P}/export.xlsx", headers=ajeno, params={"year": 2026}).status_code == 403


def test_excel_three_sheets_match_dashboard(client):
    admin = login(client, "admin@x.org")
    put_target(client, admin, "v3", 100)
    put_target(client, admin, "v4", 40, gestor="g1")
    put_target(client, admin, "v4", 40, gestor="g1", municipality="Turbo", unit="families")
    report = dashboard(client, admin, month=11)
    response = client.get(f"/api/v1/planner/{P}/export.xlsx", headers=admin, params={"year": 2026, "month": 11})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Resumen municipal", "Detalle por gestor", "Consolidado general"]
    summary = list(book["Resumen municipal"].iter_rows(min_row=2, values_only=True))
    assert sum(row[5] for row in summary) == report["totals"]["activities"]["all_executed"] == 168
    # El cumplimiento total solo mide lo que tiene meta: Apartadó (3, sin meta) va aparte.
    assert report["totals"]["activities"]["executed"] == 165 and report["totals"]["activities"]["unplanned"] == 3
    general = list(book["Consolidado general"].iter_rows(min_row=2, values_only=True))
    assert sum(row[5] for row in general if row[0] in {"Noviembre"} and row[3] != "Total del mes") == 168
    detail = list(book["Detalle por gestor"].iter_rows(min_row=2, values_only=True))
    assert sum(row[10] for row in detail) == 168  # el detalle cuadra con el resumen municipal
    v4 = municipal(report, "Visita 4 VMC")
    assert v4["activities"]["target"] == 40  # sin meta municipal, suma las metas de los gestores
    assert ("Sin meta" in [row[7] for row in summary])  # Apartadó


def test_plans_inherit_and_reschedule_with_reason(client):
    gestor = login(client, "g1@x.org")
    created = client.post(f"/api/v1/planner/{P}/plans", headers=gestor, json={
        "template_id": "v3", "start_date": "2026-11-01", "end_date": "2026-11-30", "planned_quantity": 90, "planned_participants": 50})
    assert created.status_code == 201, created.text
    plan = created.json()
    assert plan["municipality"] == "Turbo" and plan["coordinator_user_id"] == "coord"
    assert client.post(f"/api/v1/planner/{P}/plans", headers=gestor, json={
        "template_id": "v3", "gestor_user_id": "g2", "start_date": "2026-11-01", "end_date": "2026-11-30"}).status_code == 403
    assert client.patch(f"/api/v1/planner/plans/{plan['id']}/reschedule", headers=gestor, json={"reason": ""}).status_code == 422
    moved = client.patch(f"/api/v1/planner/plans/{plan['id']}/reschedule", headers=gestor,
                         json={"start_date": "2026-11-15", "end_date": "2026-12-15", "reason": "Lluvias en la vereda"})
    assert moved.json()["status"] == "rescheduled"
    # Diciembre no tiene meta ni ejecución, pero la actividad reprogramada aparece en el detalle.
    december = [r for r in dashboard(client, login(client, "admin@x.org"), month=12)["gestor"] if r["gestor"] == "Gestor Uno"]
    assert len(december) == 1 and december[0]["planned_quantity"] == 90 and december[0]["state"] == "Sin meta"
    assert december[0]["start_date"] == "2026-11-15" and december[0]["coordinator"] == "Coordinadora Turbo"
    other = login(client, "g2@x.org")
    assert client.post(f"/api/v1/planner/plans/{plan['id']}/cancel", headers=other, json={"reason": "No es mío"}).status_code == 404
    with client.sessions() as db:
        assert db.query(PlannerHistory).filter(PlannerHistory.entity_id == plan["id"]).count() == 2


def test_team_coordinator_is_configurable(client):
    admin = login(client, "admin@x.org")
    response = client.put("/api/v1/gestor-teams/team-1/coordinator", headers=admin, json={"coordinator_user_id": "g2"})
    assert response.status_code == 200 and response.json()["coordinator_user_id"] == "g2"
    assert client.put("/api/v1/gestor-teams/team-1/coordinator", headers=admin, json={"coordinator_user_id": "ajeno"}).status_code == 422
    assert client.put("/api/v1/gestor-teams/team-1/coordinator", headers=login(client, "g1@x.org"), json={"coordinator_user_id": None}).status_code == 403
